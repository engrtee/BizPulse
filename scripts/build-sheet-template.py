"""Builds docs/templates/bizpulse-sheet-template.xlsx — the starter Google Sheet for Sheets mode.
Import into Google Sheets (File -> Import -> Replace spreadsheet) to keep an existing link.
Usage: python scripts/build-sheet-template.py [--blank]   (--blank = headers only, no sample data)
"""
import sys
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

BLANK = '--blank' in sys.argv
ROWS = 500
NAVY, GREEN = '0F2744', '1A7A4A'
NAIRA = '"₦"#,##0'
today = date.today()

wb = Workbook()
head_fill = PatternFill('solid', fgColor=NAVY)
head_font = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
input_fill = PatternFill('solid', fgColor='FFF9E6')   # yellow = staff type here
calc_fill = PatternFill('solid', fgColor='EEF2F7')    # grey = formula, don't type

def header(ws, names, widths):
    ws.append(names)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'

# ── README ────────────────────────────────────────────────────────────────
rd = wb.active
rd.title = 'README'
lines = [
    ('BizPulse sheet — how to use', True),
    ('', False),
    ('1. When you BUY goods, add one row in the Purchases tab.', False),
    ('2. When you SELL goods, add one row in the Sales tab.', False),
    ('3. Yellow cells are for typing. Grey cells fill themselves — do not type in them.', False),
    ('4. Pick the item from the dropdown so the name is always spelled the same way.', False),
    ('5. New product? Add it as a new row in the Stock tab first (Item, Unit, Opening Stock, Reorder Level).', False),
    ('', False),
    ('Stock tab: "Current Stock" = Opening Stock + everything bought − everything sold. It updates by itself.', False),
    ('Red = finished. Orange = at or below your Reorder Level.', False),
    ('', False),
    ('Please do not rename the tabs or the column titles in row 1 — BizPulse reads them.', False),
    ('You can add extra columns (e.g. Notes) anywhere; BizPulse ignores them.', False),
    ('', False),
    ('Connect to BizPulse: share this sheet with the BizPulse email as Viewer, then send the link to Kemi on WhatsApp.', False),
]
for i, (t, bold) in enumerate(lines, 1):
    c = rd.cell(row=i, column=1, value=t)
    c.font = Font(bold=bold, size=14 if bold else 11, color=NAVY if bold else '000000')
    c.alignment = Alignment(wrap_text=True, vertical='top')
rd.column_dimensions['A'].width = 100

# ── Stock ─────────────────────────────────────────────────────────────────
st = wb.create_sheet('Stock')
header(st, ['Item', 'Unit', 'Opening Stock', 'Current Stock', 'Reorder Level', 'Unit Cost'], [28, 12, 15, 15, 15, 14])
items = [  # item, unit, opening, reorder, unit cost
    ('Peak Milk (tin)',   'tin',    60,  20, 1800),
    ('Indomie (carton)',  'carton', 40,  10, 7200),
    ('Coca-Cola 50cl',    'crate',  25,   8, 5400),
    ('Bottled Water',     'pack',   50,  15, 1100),
    ('Sugar 1kg',         'pack',   45,  12, 1600),
    ('Rice 50kg',         'bag',    12,   4, 78000),
    ('Vegetable Oil 5L',  'keg',    20,   6, 9500),
]
if BLANK:
    items = []
for r in range(2, ROWS + 2):
    k = r - 2
    if k < len(items):
        it, unit, op, ro, cost = items[k]
        st.cell(r, 1, it); st.cell(r, 2, unit); st.cell(r, 3, op); st.cell(r, 5, ro); st.cell(r, 6, cost)
    st.cell(r, 4, f'=IF(A{r}="","",C{r}+SUMIF(Purchases!$B:$B,A{r},Purchases!$C:$C)-SUMIF(Sales!$B:$B,A{r},Sales!$C:$C))')
    for c in (1, 2, 3, 5, 6):
        st.cell(r, c).fill = input_fill
    st.cell(r, 4).fill = calc_fill
    st.cell(r, 6).number_format = NAIRA
st.conditional_formatting.add(f'A2:F{ROWS+1}', FormulaRule(formula=['AND($A2<>"",$D2<=0)'], fill=PatternFill('solid', bgColor='F8C9C9'), font=Font(color='9B1C1C', bold=True)))
st.conditional_formatting.add(f'A2:F{ROWS+1}', FormulaRule(formula=['AND($A2<>"",$E2<>"",$D2>0,$D2<=$E2)'], fill=PatternFill('solid', bgColor='FDE6C8'), font=Font(color='9A5B00')))

# ── Purchases ─────────────────────────────────────────────────────────────
pu = wb.create_sheet('Purchases')
header(pu, ['Date', 'Item', 'Qty', 'Unit Cost', 'Total', 'Supplier'], [14, 28, 10, 14, 16, 22])
# ── Sales ─────────────────────────────────────────────────────────────────
sa = wb.create_sheet('Sales')
header(sa, ['Date', 'Item', 'Qty', 'Unit Price', 'Total', 'Customer'], [14, 28, 10, 14, 16, 22])

purchases, sales = [], []
if not BLANK:
    d = lambda n: today - timedelta(days=n)
    purchases = [
        (d(6), 'Peak Milk (tin)', 24, 1800, 'Alaba Wholesale'),
        (d(5), 'Indomie (carton)', 20, 7200, 'Mile 12 Depot'),
        (d(4), 'Bottled Water', 30, 1100, 'Alaba Wholesale'),
        (d(2), 'Vegetable Oil 5L', 6, 9500, 'Mile 12 Depot'),
    ]
    sales = [  # date, item, qty, price, customer
        (d(6), 'Peak Milk (tin)', 12, 2200, ''), (d(6), 'Indomie (carton)', 8, 8500, 'Mama Tobi'),
        (d(5), 'Coca-Cola 50cl', 6, 6200, ''), (d(5), 'Sugar 1kg', 15, 1900, ''),
        (d(4), 'Rice 50kg', 4, 86000, 'Bukka Ade'), (d(4), 'Bottled Water', 14, 1400, ''),
        (d(3), 'Peak Milk (tin)', 18, 2200, ''), (d(3), 'Vegetable Oil 5L', 9, 11000, 'Mama Tobi'),
        (d(2), 'Indomie (carton)', 22, 8500, 'Bukka Ade'), (d(2), 'Coca-Cola 50cl', 9, 6200, ''),
        (d(1), 'Rice 50kg', 8, 86000, 'Bukka Ade'), (d(1), 'Sugar 1kg', 18, 1900, ''),
        (d(1), 'Vegetable Oil 5L', 12, 11000, ''), (d(1), 'Bottled Water', 20, 1400, ''),
        (today, 'Peak Milk (tin)', 25, 2200, ''), (today, 'Indomie (carton)', 12, 8500, 'Mama Tobi'),
    ]

for ws, rows, money_col in ((pu, purchases, 4), (sa, sales, 4)):
    for r in range(2, ROWS + 2):
        k = r - 2
        if k < len(rows):
            row = rows[k]
            ws.cell(r, 1, row[0]); ws.cell(r, 2, row[1]); ws.cell(r, 3, row[2]); ws.cell(r, 4, row[3])
            ws.cell(r, 6, row[4] or None)
        ws.cell(r, 5, f'=IF(C{r}="","",C{r}*D{r})')
        for c in (1, 2, 3, 4, 6):
            ws.cell(r, c).fill = input_fill
        ws.cell(r, 5).fill = calc_fill
        ws.cell(r, 1).number_format = 'dd/mm/yyyy'
        ws.cell(r, 4).number_format = NAIRA
        ws.cell(r, 5).number_format = NAIRA
    dv = DataValidation(type='list', formula1=f'=Stock!$A$2:$A${ROWS+1}', allow_blank=True, showErrorMessage=True,
                        errorTitle='Unknown item', error='Pick an item from the list, or add it to the Stock tab first.')
    ws.add_data_validation(dv)
    dv.add(f'B2:B{ROWS+1}')

out = 'docs/templates/bizpulse-sheet-template' + ('-blank' if BLANK else '') + '.xlsx'
import os
os.makedirs('docs/templates', exist_ok=True)
wb.save(out)
print('wrote', out)
